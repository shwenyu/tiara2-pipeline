#!/usr/bin/env python3
"""Import Tiara's original training corpus (Supplementary Table S1) into our
corpus layout, deduplicated against whatever we already hold.

Tiara's published set is 8,198 assemblies: 3,821 prokaryotic chromosomes,
2,260 mitochondria, 2,048 plastids and 69 eukaryotic nuclear genomes.  Those
69 are worth more than their count suggests -- they were hand-picked, and they
still beat our much larger euk set -- so they are treated as anchors and are
never dropped by a downstream quality gate.

What this does
--------------
1. Normalises every accession in S1 (regex + version stripping), the same way
   ``inspect_external_metadata.py`` does.
2. Diffs them against our canonical index: exact hits, base-accession hits
   (same assembly, different version), and genuinely new accessions.
3. Assigns each row a fine class, a stage-1 class and -- for eukaryotes -- a
   stage-2 clade, using tiara2.labels so the import cannot disagree with the
   rest of the pipeline.
4. Writes a download plan laid out the way we store things:
   ``{base}/raw/{group}/{split}/`` , plus a selection TSV shaped like the
   ``curate`` output so the existing stages can consume it unchanged.

Usage
-----
    python3 scripts/import_tiara_corpus.py \\
        --external Supplementary_Table_S1.xlsx \\
        --index /data/shouhanyu/Tiara2/index/all_entities.tsv \\
        --out-dir /data/shouhanyu/Tiara2/import/tiara1

    # regenerate the anchor list consumed by curate
    python3 scripts/import_tiara_corpus.py --external ... --out-dir ... \\
        --emit-anchor-file config/tiara1_anchor_accessions.txt
"""

from __future__ import annotations

import argparse
import csv
import gzip
import json
import os
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

sys.path.insert(0, str(Path(__file__).resolve().parent))

from progress import (  # noqa: E402
    Counter as ProgressCounter, add_progress_args, counter_from_args,
)
from tiara2 import labels as L  # noqa: E402

# --- accession handling (shared shape with inspect_external_metadata.py) ----
ACCESSION_RE = re.compile(
    r"(?<![A-Z0-9])(?:"
    r"GC[AF]_\d+(?:\.\d+)?|"
    r"(?:NC|NW|NZ|NT|NG|NM|NR|XM|XR|NP|YP|WP)_\d+(?:\.\d+)?|"
    r"[A-Z]{1,4}\d{5,}(?:\.\d+)?"
    r")(?![A-Z0-9])",
    re.I,
)
DEFAULT_INDEX_COLUMNS = ("entity_id", "assembly_accession", "sequence_accession",
                         "paired_accession", "source_native_id")


def norm(value: object) -> str:
    return str(value or "").strip().upper()


def base_accession(acc: str) -> str:
    return re.sub(r"\.\d+$", "", norm(acc))


def extract_accessions(value: object) -> list[str]:
    text = norm(value)
    if not text:
        return []
    found = [norm(x) for x in ACCESSION_RE.findall(text)]
    if not found and (text.startswith("JGI:") or text.startswith("IMG:")):
        found = [text]
    return list(dict.fromkeys(found))


def open_text(path: str):
    if str(path).lower().endswith(".gz"):
        return gzip.open(path, "rt", encoding="utf-8", errors="replace", newline="")
    return open(path, encoding="utf-8", errors="replace", newline="")


# --- S1 vocabulary ---------------------------------------------------------
# Exact strings from Supplementary Table S1; matched case-insensitively and
# exactly -- never by substring.
GENOME_TYPE_TO_CLASS = {
    "mitochondrion": "mitochondria",
    "plastid": "plastids",
    "eukarya nuclear": "eukarya",
}
DOMAIN_TO_CLASS = {
    "archaea": "archaea",
    "bacteria": "bacteria",
    "eukarya": "eukarya",
}
# Tiara's `Taxonomy` column for the 69 nuclear eukaryotes -> our stage-2 clade
# and the NCBI division we would download it from.
TAXONOMY_TO_CLADE = {
    "fungi": ("fungi", "fungi"),
    "ascomycota": ("fungi", "fungi"),
    "basidiomycota": ("fungi", "fungi"),
    "microsporidia": ("fungi", "fungi"),
    "alveolata": ("alveolata", "protozoa"),
    "apicomplexa": ("alveolata", "protozoa"),
    "ciliophora": ("alveolata", "protozoa"),
    "stramenopiles": ("stramenopiles", "protozoa"),
    "bacillariophyta": ("stramenopiles", "protozoa"),
    "oomycota": ("stramenopiles", "protozoa"),
    "euglenozoa": ("other_protist", "protozoa"),
    "kinetoplastida": ("other_protist", "protozoa"),
    "discoba": ("other_protist", "protozoa"),
    "evosea": ("other_protist", "protozoa"),
    "amoebozoa": ("other_protist", "protozoa"),
    "rhizaria": ("other_protist", "protozoa"),
    "chlorophyta": ("algae", "plant"),
    "rhodophyta": ("algae", "plant"),
    "bangiophyceae": ("algae", "plant"),
    "glaucocystophyceae": ("algae", "plant"),
    "streptophyta": ("land_plant", "plant"),
    "embryophyta": ("land_plant", "plant"),
    "viridiplantae": ("land_plant", "plant"),
    "metazoa": ("metazoa_invertebrate", "invertebrate"),
    "arthropoda": ("metazoa_invertebrate", "invertebrate"),
    "nematoda": ("metazoa_invertebrate", "invertebrate"),
    "chordata": ("metazoa_vertebrate", "vertebrate_other"),
    "vertebrata": ("metazoa_vertebrate", "vertebrate_other"),
    "mammalia": ("metazoa_vertebrate", "vertebrate_mammalian"),
}
# Where each fine class is stored under {base}/raw/.
GROUP_BY_CLASS = {
    "bacteria": "bacteria",
    "archaea": "archaea",
    "mitochondria": "mitochondrion",
    "plastids": "plastid",
    "virus": "viral",
}


def classify_row(genome_type: str, domain: str, taxonomy: str) -> dict:
    """Map one S1 row onto our class hierarchy."""
    gt = (genome_type or "").strip().lower()
    dm = (domain or "").strip().lower()

    fine = GENOME_TYPE_TO_CLASS.get(gt)
    if fine is None and gt == "prokaryotic chromosome":
        fine = DOMAIN_TO_CLASS.get(dm)
    if fine is None:
        fine = DOMAIN_TO_CLASS.get(dm)
    if fine is None:
        return {"fine_class": "", "stage1_class": "", "stage2_clade": "",
                "clade_source": "unmapped", "group": ""}

    stage1 = L.stage1_of(fine) or ""
    clade, how, group = "", "", GROUP_BY_CLASS.get(fine, "")
    if fine == "eukarya":
        tokens = [t.strip().lower() for t in re.split(r"[;,/|>]+", taxonomy or "") if t.strip()]
        for token in reversed(tokens):        # most specific rank first
            if token in TAXONOMY_TO_CLADE:
                clade, group = TAXONOMY_TO_CLADE[token]
                how = "taxonomy_string"
                break
        if not clade:
            clade, how, group = "eukarya_other", "fallback", "protozoa"
    return {"fine_class": fine, "stage1_class": stage1, "stage2_clade": clade,
            "clade_source": how, "group": group}


# --- readers ---------------------------------------------------------------
def iter_xlsx(path: str, sheet: str | None, header_row: int | None):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError("reading XLSX needs openpyxl: pip install openpyxl") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    headers = None
    if header_row:
        for i, row in enumerate(rows, 1):
            if i == header_row:
                headers = [str(x or "").strip() for x in row]
                break
    else:
        for i, row in enumerate(rows, 1):
            vals = [str(x or "").strip() for x in row]
            if any("accession" in v.lower() for v in vals):
                headers = vals
                break
            if i >= 30:
                break
    if headers is None:
        raise ValueError("no header row containing 'Accession' in the first 30 rows")
    for rn, row in enumerate(rows, 1):
        yield rn, {headers[i] or f"column_{i+1}": (row[i] if i < len(row) else "")
                   for i in range(len(headers))}


def iter_delimited(path: str):
    with open_text(path) as handle:
        sample = handle.read(65536)
        handle.seek(0)
        ext = Path(path).suffix.lower()
        delim = "\t" if ext == ".tsv" else "," if ext == ".csv" else None
        if delim is None:
            try:
                delim = csv.Sniffer().sniff(sample, delimiters="\t,;").delimiter
            except csv.Error:
                delim = "\t"
        for rn, row in enumerate(csv.DictReader(handle, delimiter=delim), 2):
            yield rn, row


def iter_external(path: str, sheet: str | None, header_row: int | None):
    if str(path).lower().endswith((".xlsx", ".xlsm")):
        return iter_xlsx(path, sheet, header_row)
    return iter_delimited(path)


def load_index(path: str | None, columns=DEFAULT_INDEX_COLUMNS, prog=None):
    """Return ``(exact, base)`` accession -> list of index hits."""
    prog = prog or ProgressCounter("import", enabled=False)
    exact: dict[str, list] = defaultdict(list)
    base: dict[str, list] = defaultdict(list)
    if not path:
        return exact, base
    prog.write(f"index: reading {path}")
    with open_text(path) as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        present = [c for c in columns if c in (reader.fieldnames or [])]
        if not present:
            raise ValueError(f"index has no accession column; available={reader.fieldnames}")
        for rn, row in enumerate(reader, 2):
            prog.count()
            for col in present:
                for acc in extract_accessions(row.get(col, "")):
                    hit = {"index_row": rn, "index_column": col, "index_accession": acc,
                           "index_entity_id": row.get("entity_id", "")}
                    exact[acc].append(hit)
                    base[base_accession(acc)].append(hit)
    return exact, base


# --- core ------------------------------------------------------------------
COLUMNS = ["accession", "base_accession", "organism_name", "taxonomy", "genome_type",
           "domain", "fine_class", "stage1_class", "stage2_clade", "clade_source",
           "group", "genome_length_mb", "gc_percent", "status", "is_anchor",
           "target_dir", "source_row"]


def build_rows(paths, sheet=None, header_row=None, raw_root="", split="train",
               prog=None):
    rows, invalid = [], []
    # Row-oriented input: openpyxl/csv expose no byte offset, so this reports
    # rows and leaves percent/ETA as '?' rather than inventing them.
    prog = prog or ProgressCounter("import", enabled=False)
    for path in paths:
        prog.write(f"external: reading {path}")
        for rn, record in iter_external(path, sheet, header_row):
            prog.count()
            def get(*names):
                for name in names:
                    if name in record and str(record[name] or "").strip():
                        return str(record[name]).strip()
                return ""
            accessions = extract_accessions(get("Accession", "accession"))
            if not accessions:
                invalid.append({"source_file": os.path.abspath(str(path)), "source_row": rn,
                                "raw_accession": get("Accession", "accession"),
                                "reason": "NO_RECOGNIZED_ACCESSION"})
                continue
            taxonomy = get("Taxonomy", "taxonomy")
            assigned = classify_row(get("Genome type", "genome_type"),
                                    get("Domain", "domain"), taxonomy)
            if not assigned["fine_class"]:
                invalid.append({"source_file": os.path.abspath(str(path)), "source_row": rn,
                                "raw_accession": accessions[0], "reason": "UNMAPPED_CLASS"})
                continue
            for acc in accessions:
                target = ""
                if raw_root and assigned["group"]:
                    target = str(Path(raw_root) / assigned["group"] / split)
                rows.append({
                    "accession": acc,
                    "base_accession": base_accession(acc),
                    "organism_name": get("Organism name", "organism_name"),
                    "taxonomy": taxonomy,
                    "genome_type": get("Genome type", "genome_type"),
                    "domain": get("Domain", "domain"),
                    "genome_length_mb": get("Genome Length (Mb)"),
                    "gc_percent": get("Genome GC content (%)"),
                    "status": "",
                    "is_anchor": "1" if assigned["fine_class"] == "eukarya" else "0",
                    "target_dir": target,
                    "source_row": rn,
                    **assigned,
                })
        # Per-file verdict, immediately: a wrong sheet or header row shows up
        # as "0 usable rows" here instead of in a silent empty report.
        prog.write(f"external: {path} -> {len(rows):,} row(s) so far, "
                   f"{len(invalid):,} unusable")
    return rows, invalid


def diff_against_index(rows, exact, base):
    seen_exact: dict[str, int] = {}
    seen_base: dict[str, int] = {}
    for row in rows:
        acc, bacc = row["accession"], row["base_accession"]
        if acc in seen_exact:
            row["status"] = "duplicate_in_source"
        elif bacc in seen_base:
            row["status"] = "duplicate_in_source_other_version"
        elif acc in exact:
            row["status"] = "already_held"
        elif bacc in base:
            row["status"] = "already_held_other_version"
        else:
            row["status"] = "new"
        seen_exact.setdefault(acc, 1)
        seen_base.setdefault(bacc, 1)
    return rows


def summarise(rows, invalid):
    by_status = Counter(r["status"] for r in rows)
    by_stage1 = Counter(r["stage1_class"] for r in rows)
    by_fine = Counter(r["fine_class"] for r in rows)
    by_clade = Counter(r["stage2_clade"] for r in rows if r["stage2_clade"])
    new_by_fine = Counter(r["fine_class"] for r in rows if r["status"] == "new")
    return {
        "rows": len(rows),
        "unique_accessions": len({r["accession"] for r in rows}),
        "unique_base_accessions": len({r["base_accession"] for r in rows}),
        "invalid_rows": len(invalid),
        "by_status": dict(by_status),
        "by_stage1": dict(by_stage1),
        "by_fine_class": dict(by_fine),
        "by_stage2_clade": dict(by_clade),
        "new_by_fine_class": dict(new_by_fine),
        "anchors": sum(1 for r in rows if r["is_anchor"] == "1"),
        "clade_source": dict(Counter(r["clade_source"] for r in rows if r["stage2_clade"])),
    }


def write_tsv(path, rows, columns):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, delimiter="\t",
                                extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def render(summary) -> str:
    lines = ["tiara1 corpus import", ""]
    lines.append(f"  rows normalised     {summary['rows']:>8,}")
    lines.append(f"  unique accessions   {summary['unique_accessions']:>8,}")
    lines.append(f"  unmappable rows     {summary['invalid_rows']:>8,}")
    lines.append("")
    lines.append("  stage 1:")
    for name in L.STAGE1_CLASSES:
        lines.append(f"      {name:22} {summary['by_stage1'].get(name, 0):>8,}")
    lines.append("  fine classes:")
    for name, n in sorted(summary["by_fine_class"].items(), key=lambda kv: -kv[1]):
        new = summary["new_by_fine_class"].get(name, 0)
        lines.append(f"      {name:22} {n:>8,}   new: {new:>8,}")
    if summary["by_stage2_clade"]:
        lines.append("  stage 2 (eukarya only):")
        for name in L.STAGE2_EUK_CLASSES:
            n = summary["by_stage2_clade"].get(name, 0)
            if n:
                lines.append(f"      {name:22} {n:>8,}")
    lines.append("")
    lines.append("  vs our index:")
    for status, n in sorted(summary["by_status"].items(), key=lambda kv: -kv[1]):
        lines.append(f"      {status:34} {n:>8,}")
    lines.append("")
    lines.append(f"  anchors (euk nuclear) {summary['anchors']:>6,}")
    return "\n".join(lines)


# Where a Supplementary table plausibly sits on this box. Used ONLY to make a
# missing --external path actionable: openpyxl's own failure for a nonexistent
# file is a 15-frame traceback ending in zipfile, which says nothing about which
# argument was wrong.
SEARCH_ROOTS = (
    ".", "metadata", "docs", "config",
    "/data/shouhanyu/Tiara2/metadata", "/data/shouhanyu/Tiara2", "/data/shouhanyu",
    str(Path.home()), str(Path.home() / "tiara2_pipeline"),
)
SEARCH_GLOBS = ("*upplementary*.xlsx", "*upplementary*.csv", "*upplementary*.tsv",
                "S1*.xlsx", "*able_S1*", "*table_s1*")


def _load_cfg(config_path: str) -> dict:
    """Load config.yaml, tolerating a cwd outside the checkout."""
    from tiara2 import config as config_mod
    from tiara2 import paths as paths_mod

    path = Path(os.path.expanduser(str(config_path)))
    if not path.exists() and not path.is_absolute():
        path = paths_mod.repo_root() / config_path
    return config_mod.load(path)


def default_out_dir(config_path: str) -> str:
    """``{base}/import/tiara1`` -- the COLD tier.

    Tiering is decided by READ COUNT, not size. These TSVs are written once and
    read once (by eyes, or by a download step), so NVMe buys nothing; the hot
    tier is for what every k and every HP candidate re-reads.
    """
    try:
        cfg = _load_cfg(config_path)
    except Exception as exc:                      # noqa: BLE001 - config is optional here
        raise SystemExit(f"cannot resolve a default --out-dir ({exc});"
                         " pass --out-dir explicitly")
    return str(Path(str(cfg.get("base") or ".")) / "import" / "tiara1")


def warn_hot_tier(out_dir: Path, config_path: str) -> None:
    """Say something when plan artifacts are being written to the hot tier.

    Not an error: a single-disk setup has fast_base == base, and someone may
    have a reason. But silently spending NVMe on write-once files is exactly
    the mistake the two-tier layout exists to avoid.
    """
    try:
        cfg = _load_cfg(config_path)
    except Exception:                             # noqa: BLE001
        return
    base, fast_base = str(cfg.get("base") or ""), str(cfg.get("fast_base") or "")
    if not fast_base or fast_base == base:
        return
    resolved = str(Path(out_dir).expanduser())
    if resolved.startswith(fast_base):
        print(f"warning: {out_dir} is on the HOT tier ({fast_base}).\n"
              f"         These artifacts are written once and read once; "
              f"{Path(base) / 'import' / 'tiara1'} is the right place.",
              file=sys.stderr)


def find_candidates(limit: int = 12) -> list[Path]:
    """Best-effort search for Supplementary tables, nearest roots first."""
    seen: dict[str, Path] = {}
    for root in SEARCH_ROOTS:
        base_dir = Path(root)
        if not base_dir.is_dir():
            continue
        for pattern in SEARCH_GLOBS:
            try:
                matches = sorted(base_dir.glob(pattern))
            except OSError:
                continue
            for match in matches:
                if match.is_file():
                    seen.setdefault(str(match.resolve()), match)
                if len(seen) >= limit:
                    return list(seen.values())
    return list(seen.values())


def check_external(paths) -> list[str]:
    """Return a human-readable error list for any --external path we cannot read."""
    errors = []
    for raw in paths:
        path = Path(os.path.expanduser(str(raw)))
        if path.is_dir():
            errors.append(f"--external {raw} is a directory, not a table file")
        elif not path.exists():
            errors.append(f"--external {raw} does not exist (cwd: {Path.cwd()})")
        elif path.stat().st_size == 0:
            errors.append(f"--external {raw} is empty")
    return errors


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--external", nargs="+", required=True,
                    help="Supplementary_Table_S1.xlsx (or a TSV/CSV export)")
    ap.add_argument("--index", default="", help="our canonical all_entities.tsv")
    ap.add_argument("--out-dir", default="",
                    help="where the TSV artifacts go; default {base}/import/tiara1 "
                         "on the COLD tier (these are write-once, read-once "
                         "plan artifacts -- NVMe buys nothing)")
    ap.add_argument("--config", default="config/config.yaml",
                    help="only used to resolve the default --out-dir")
    ap.add_argument("--raw-root", default="",
                    help="root the download plan targets, e.g. {base}/raw")
    ap.add_argument("--split", default="train")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--header-row", type=int, default=None)
    ap.add_argument("--emit-anchor-file", default="",
                    help="also write the euk nuclear accessions as a curate anchor file")
    add_progress_args(ap)
    args = ap.parse_args(argv)

    # Fail on the ARGUMENT, not 15 frames deep inside openpyxl/zipfile.
    errors = check_external(args.external)
    if errors:
        for message in errors:
            print(f"error: {message}", file=sys.stderr)
        candidates = find_candidates()
        if candidates:
            print("\nfound these Supplementary-looking tables instead:", file=sys.stderr)
            for candidate in candidates:
                print(f"  {candidate}", file=sys.stderr)
            print("\nRe-run with an absolute path to the right one.", file=sys.stderr)
        else:
            print("\nNo Supplementary table found under: "
                  + ", ".join(SEARCH_ROOTS), file=sys.stderr)
        return 2

    out_dir = Path(args.out_dir) if args.out_dir else Path(default_out_dir(args.config))
    if not args.out_dir:
        print(f"out-dir defaulted to {out_dir}", file=sys.stderr)
    warn_hot_tier(out_dir, args.config)
    out_dir.mkdir(parents=True, exist_ok=True)

    # Progress lines go to stderr; the rendered summary below is the result and
    # stays on stdout.
    prog = counter_from_args(args, label="import")
    rows, invalid = build_rows(args.external, args.sheet, args.header_row,
                               args.raw_root, args.split, prog=prog)
    exact, base = load_index(args.index or None, prog=prog)
    prog.write(f"diffing {len(rows):,} row(s) against the index")
    rows = diff_against_index(rows, exact, base)
    summary = summarise(rows, invalid)
    prog.finish(f"{len(rows):,} selected, {len(invalid):,} unusable")

    write_tsv(out_dir / "tiara1_selection.tsv", rows, COLUMNS)
    write_tsv(out_dir / "new_accessions.tsv",
              [r for r in rows if r["status"] == "new"], COLUMNS)
    write_tsv(out_dir / "already_held.tsv",
              [r for r in rows if r["status"].startswith("already_held")], COLUMNS)
    write_tsv(out_dir / "duplicates_in_source.tsv",
              [r for r in rows if r["status"].startswith("duplicate")], COLUMNS)
    write_tsv(out_dir / "invalid_rows.tsv", invalid,
              ["source_file", "source_row", "raw_accession", "reason"])
    write_tsv(out_dir / "download_plan.tsv",
              [r for r in rows if r["status"] == "new" and r["target_dir"]],
              ["accession", "fine_class", "stage1_class", "stage2_clade", "group",
               "target_dir", "organism_name"])
    (out_dir / "import_report.json").write_text(
        json.dumps({"summary": summary, "external": [os.path.abspath(str(p)) for p in args.external],
                    "index": args.index}, indent=2, ensure_ascii=False))

    if args.emit_anchor_file:
        anchors = sorted({r["accession"] for r in rows if r["is_anchor"] == "1"})
        path = Path(args.emit_anchor_file)
        path.parent.mkdir(parents=True, exist_ok=True)
        header = [
            "# Tiara1 eukaryotic nuclear anchor assemblies, generated by",
            "# scripts/import_tiara_corpus.py from Supplementary Table S1.",
            "# These are never dropped by a curate quality gate.",
            f"# count: {len(anchors)}",
        ]
        path.write_text("\n".join(header + anchors) + "\n")
        print(f"wrote {len(anchors)} anchors -> {path}")

    print(render(summary))
    print(f"\nartifacts -> {out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
